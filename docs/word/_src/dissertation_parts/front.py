"""Title page, approval page, abstract, acknowledgments, lists, abbreviations."""
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

TITLE = ("Fully Offline, Layout-Preserving Image Translation on Commodity Smartphones: "
         "A Dependency-Light Pipeline with Translucent Re-Rendering")


def _centered(c, text, size=12, bold=False, italic=False, space_after=6):
    para = c.d.d.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold or None
    run.italic = italic or None
    para.paragraph_format.space_after = Pt(space_after)
    return para


def build(c):
    d = c.d
    # ------------------------------------------------------------ title page
    for _ in range(3):
        d.d.add_paragraph()
    _centered(c, TITLE, size=18, bold=True, space_after=24)
    _centered(c, "by", size=12, italic=True)
    _centered(c, "[Author Name]", size=14, bold=True, space_after=30)
    _centered(c, "A Dissertation Submitted in Partial Fulfillment of the Requirements "
                 "for the Degree of Doctor of Philosophy", size=12, space_after=18)
    _centered(c, "[Department]", size=12)
    _centered(c, "[University]", size=12, space_after=30)
    _centered(c, "Advisor: [Advisor]", size=12, space_after=30)
    _centered(c, "September 2026", size=12)
    d.page_break()

    # --------------------------------------------------------- approval page
    d.h(1, "Approval Page")
    c.p("This dissertation, entitled *" + TITLE + "*, submitted by [Author Name] in partial "
        "fulfillment of the requirements for the degree of Doctor of Philosophy in [Department], "
        "has been examined and approved by the committee listed below.", align="justify")
    d.table(["Role", "Name", "Signature", "Date"],
            [["Advisor", "[Advisor]", "", ""],
             ["Committee member", "[Committee Member 1]", "", ""],
             ["Committee member", "[Committee Member 2]", "", ""],
             ["Committee member", "[Committee Member 3]", "", ""],
             ["External examiner", "[External Examiner]", "", ""],
             ["Head of department", "[Department Head]", "", ""]],
            widths_cm=[4, 5, 4.5, 2.8], caption=None, font_size=10)
    c.p("[Placeholder: the approval page is completed by the graduate school according to its "
        "own template.]", italic=True, color=None)
    d.page_break()

    # --------------------------------------------------------------- abstract
    d.h(1, "Abstract")
    c.ps(
        "Image translation applications recognise the text in a photograph or scanned page, translate "
        "it, and paint the translation back over the original. Most deployed systems send the image to "
        "a server. That design is convenient for the provider, but it excludes the documents for which "
        "translation is often most needed — medical and legal paperwork, identity documents, material "
        "under data-residency rules — and it fails wherever there is no network connection. This "
        "dissertation designs, analyses and specifies the evaluation of a complete image translation "
        "system for Android that runs entirely on the device, requests no network permission, and is "
        "written in managed Java with no native code of its own beyond the inference runtime.",
        "The system chains PP-OCR text detection, 180° orientation classification and connectionist "
        "temporal classification (CTC) recognition with OPUS-MT Transformer translation; all models are "
        "executed by ONNX Runtime, the translation models on dynamically quantised int8 weights. Three "
        "technical contributions make the pipeline work within the constraints of a phone. First, the "
        "Differentiable Binarization (DB) post-processing that reference implementations delegate to "
        "OpenCV and a polygon-clipping library is reimplemented from first principles. We prove that "
        "restricting the convex-hull input to the leftmost and rightmost pixel of each row is exact "
        "(Lemma 4.1), derive a closed-form unclip offset for rectangles, and score candidate regions "
        "over their own pixels rather than their bounding boxes to remove a systematic bias against "
        "rotated text. Second, translation between Chinese, Japanese and Korean, for which no compact "
        "direct models are published, is routed through English with whole-page batching per leg and "
        "at most one model pair resident in memory; the cost of a decoder without a key-value cache is "
        "analysed exactly as n(n+1)/2 decoder positions for n generated tokens. Third, the usual opaque "
        "erase of the original text is replaced by a translucent, feathered highlight in the locally "
        "estimated paper colour, combined with a paper-coloured glyph halo. We prove that inside the "
        "highlight the original glyph contrast and the page texture are both attenuated by exactly "
        "(1 − α) in the compositing colour space (Proposition 6.1), which is 0.41 at the default "
        "α = 150/255, while the halo restores full contrast around every new glyph. A bisection search "
        "chooses the largest feasible font size under high-quality line breaking in a bounded number of "
        "layouts.",
        "The engine implements English, Korean, Japanese and Chinese; the deployed user interface "
        "exposes three of them through a single build-time list, and uses large, screen-relative touch "
        "targets and a collapsible overlay side menu for one-handed use. The dissertation specifies a "
        "full evaluation protocol for five research questions: faithfulness of the library-free "
        "detection post-processing against the reference implementation on ICDAR 2015 and MLT-2019; "
        "recognition accuracy with and without orientation correction; translation quality on "
        "FLORES-200 for all twelve directions with BLEU, chrF and COMET, including int8 versus fp32, "
        "greedy versus beam search, and pivoted versus direct translation; on-device latency and memory "
        "on three device tiers; and a within-subjects user study of rendering variants analysed with "
        "mixed-effects models. The system has not yet been benchmarked. All experimental results are "
        "therefore reported as explicit placeholders pending execution of the protocol, and every "
        "number in this document is either a repository fact or an analytical consequence of a stated "
        "formula.",
    )
    c.p("**Keywords:** optical character recognition; machine translation; on-device inference; "
        "scene text; layout preservation; image compositing; mobile computing; privacy; "
        "human–computer interaction.", align="left")
    d.page_break()

    # -------------------------------------------------------- acknowledgments
    d.h(1, "Acknowledgments")
    c.p("[Placeholder for acknowledgments: advisor, committee, colleagues, funding bodies and family.]",
        italic=True)
    c.p("The system described here uses third-party model weights: PaddleOCR models (Apache License 2.0) "
        "and OPUS-MT models by the Helsinki-NLP group (CC-BY 4.0). Their authors' work is gratefully "
        "acknowledged; attribution is also shown in the application's About dialog.")
    d.page_break()

    # ---------------------------------------------------------------- lists
    d.toc(title="Table of Contents", levels="1-2")
    d.list_of("Table")
    d.list_of("Figure")

    d.h(1, "List of Abbreviations")
    rows = [
        ("ANR", "Application Not Responding (Android watchdog)"), ("API", "Application programming interface"),
        ("APK", "Android application package"), ("ARGB", "Alpha–red–green–blue pixel format"),
        ("BLEU", "Bilingual evaluation understudy"), ("CJK", "Chinese, Japanese and Korean"),
        ("CLMM", "Cumulative link mixed model"), ("COMET", "Crosslingual optimized metric for evaluation of translation"),
        ("CRNN", "Convolutional recurrent neural network"), ("CTC", "Connectionist temporal classification"),
        ("DB", "Differentiable Binarization"), ("dp", "Density-independent pixel (Android)"),
        ("DOCX", "Office Open XML word-processing document"), ("dpi", "Dots per inch"),
        ("fp32", "32-bit floating point"), ("H-mean", "Harmonic mean of precision and recall"),
        ("HCI", "Human–computer interaction"), ("int8", "8-bit integer (quantised) arithmetic"),
        ("IoU", "Intersection over union"), ("JVM", "Java virtual machine"),
        ("KV cache", "Key–value cache of decoder self-attention states"), ("LID", "Language identification"),
        ("MT", "Machine translation"), ("NDK", "Android Native Development Kit"),
        ("NED", "Normalised edit distance"), ("NMT", "Neural machine translation"),
        ("OCR", "Optical character recognition"), ("ONNX", "Open Neural Network Exchange"),
        ("OOXML", "Office Open XML"), ("ORT", "ONNX Runtime"),
        ("PDF", "Portable Document Format"), ("PP-OCR", "PaddlePaddle OCR system family"),
        ("PSS", "Proportional set size (memory)"), ("RQ", "Research question"),
        ("sp", "Scale-independent pixel (Android text size)"), ("SPM", "SentencePiece model"),
        ("sRGB", "Standard RGB colour space"), ("TOST", "Two one-sided tests (equivalence)"),
        ("TTS", "Text to speech"), ("TXT", "Plain text"), ("UI", "User interface"),
    ]
    d.table(["Abbreviation", "Meaning"], rows, widths_cm=[3.5, 12.5], caption=None, font_size=10)
