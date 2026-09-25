"""Small house-style layer over python-docx for the project's Word documents.

Every document generator in this folder builds a :class:`Doc`, adds content
through its methods, and saves. Keeping styling here means the SRS, design,
screen and test documents look like one set.

Inline markup accepted by :meth:`Doc.p` and table cells:
    **bold**   *italic*   `code`
"""
from __future__ import annotations

import re
from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ACCENT = RGBColor(0xA8, 0x15, 0x15)      # brand_red_dark
INK = RGBColor(0x1F, 0x23, 0x2B)
MUTED = RGBColor(0x5F, 0x67, 0x78)
HEADER_FILL = "E9ECF2"
NOTE_FILL = "FFF6E5"
CODE_FILL = "F3F4F6"

_INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*\s][^*]*?\*|`[^`]+`)")


def _set_cell_fill(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def _set_paragraph_fill(paragraph, hex_fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    p_pr.append(shd)


def _add_field(run, instr: str) -> None:
    """Appends a complex field (PAGE, NUMPAGES, TOC ...) to a run."""
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instr
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "1" if "PAGE" in instr else "Right-click and choose Update Field."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (begin, text, sep, placeholder, end):
        run._r.append(el)


def _repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def _keep_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


class Doc:
    """One Word document in the project house style."""

    def __init__(self, *, title: str, subtitle: str = "", doc_id: str = "",
                 version: str = "1.0", status: str = "Draft", date: str = "",
                 owner: str = "", body_font: str = "Calibri", body_size: float = 10.5,
                 line_spacing: float = 1.15, heading_font: str | None = None,
                 margins_cm: tuple[float, float, float, float] = (2.5, 2.2, 2.5, 2.2)):
        self.title = title
        self.subtitle = subtitle
        self.doc_id = doc_id
        self.version = version
        self.status = status
        self.date = date
        self.owner = owner
        self._table_no = 0
        self._figure_no = 0
        self.chapter_prefix = ""   # e.g. "3." gives "Table 3.1"
        self._chapter_tables = 0
        self._chapter_figures = 0

        self.d = Document()
        sec = self.d.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)          # A4
        top, right, bottom, left = margins_cm
        sec.top_margin, sec.right_margin = Cm(top), Cm(right)
        sec.bottom_margin, sec.left_margin = Cm(bottom), Cm(left)
        sec.different_first_page_header_footer = True

        heading_font = heading_font or body_font
        styles = self.d.styles
        normal = styles["Normal"]
        normal.font.name = body_font
        normal.font.size = Pt(body_size)
        normal.font.color.rgb = INK
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), body_font)
        pf = normal.paragraph_format
        pf.space_after = Pt(6)
        pf.line_spacing = line_spacing

        for level, size in ((1, 17), (2, 13.5), (3, 11.5), (4, 10.5)):
            st = styles[f"Heading {level}"]
            st.font.name = heading_font
            st.element.rPr.rFonts.set(qn("w:eastAsia"), heading_font)
            st.font.size = Pt(size)
            st.font.bold = True
            st.font.italic = False
            st.font.color.rgb = ACCENT if level == 1 else INK
            st.paragraph_format.space_before = Pt(18 if level == 1 else 12)
            st.paragraph_format.space_after = Pt(6)
            st.paragraph_format.keep_with_next = True

        for name in ("List Bullet", "List Bullet 2"):
            styles[name].font.name = body_font
            styles[name].font.size = Pt(body_size)

        cap = styles["Caption"]
        cap.font.name = body_font
        cap.font.size = Pt(9)
        cap.font.bold = True
        cap.font.italic = False
        cap.font.color.rgb = MUTED

        # Ask Word to refresh the TOC and page fields when the file is opened.
        # CT_Settings is a strict sequence, so it must sit before these siblings.
        settings = self.d.settings.element
        upd = OxmlElement("w:updateFields")
        upd.set(qn("w:val"), "true")
        later = {"hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids",
                 "mathPr", "attachedSchema", "themeFontLang", "clrSchemeMapping",
                 "doNotIncludeSubdocsInStats", "doNotAutoCompressPictures", "forceUpgrade",
                 "captions", "readModeInkLockDown", "smartTagType", "schemaLibrary",
                 "shapeDefaults", "doNotEmbedSmartTags", "decimalSymbol", "listSeparator"}
        anchor = next((c for c in settings if c.tag.split("}")[-1] in later), None)
        if anchor is not None:
            anchor.addprevious(upd)
        else:
            settings.append(upd)
        # The stock template's <w:zoom> lacks the schema-required percent attribute.
        for zoom in settings.findall(qn("w:zoom")):
            if zoom.get(qn("w:percent")) is None:
                zoom.set(qn("w:percent"), "100")

        self._header_footer(sec)

    # ------------------------------------------------------------ page chrome

    def _header_footer(self, sec) -> None:
        hp = sec.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = hp.add_run(f"{self.title}" + (f"  |  {self.doc_id} v{self.version}" if self.doc_id else ""))
        r.font.size = Pt(8)
        r.font.color.rgb = MUTED

        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for part in ("Page ", None, " of ", "NUM"):
            run = fp.add_run()
            run.font.size = Pt(8)
            run.font.color.rgb = MUTED
            if part is None:
                _add_field(run, "PAGE")
            elif part == "NUM":
                _add_field(run, "NUMPAGES")
            else:
                run.text = part

    # ------------------------------------------------------------ front matter

    def cover(self, extra_rows: Sequence[tuple[str, str]] = ()) -> None:
        for _ in range(6):
            self.d.add_paragraph()
        p = self.d.add_paragraph()
        r = p.add_run(self.title)
        r.font.size = Pt(28)
        r.font.bold = True
        r.font.color.rgb = ACCENT
        if self.subtitle:
            p = self.d.add_paragraph()
            r = p.add_run(self.subtitle)
            r.font.size = Pt(14)
            r.font.color.rgb = MUTED
        rule = self.d.add_paragraph()
        self._bottom_border(rule, "A81515", 12)
        self.d.add_paragraph()

        rows = [("Document ID", self.doc_id), ("Version", self.version),
                ("Status", self.status), ("Date", self.date)]
        if self.owner:
            rows.append(("Owner", self.owner))
        rows.extend(extra_rows)
        t = self.d.add_table(rows=0, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        for k, v in rows:
            if not v:
                continue
            cells = t.add_row().cells
            cells[0].width, cells[1].width = Cm(4), Cm(10)
            self._fill(cells[0].paragraphs[0], k, bold=True, color=MUTED)
            self._fill(cells[1].paragraphs[0], v)
        self.page_break()

    def revision_history(self, rows: Sequence[Sequence[str]]) -> None:
        self.h(1, "Revision History", numbered=False)
        self.table(["Version", "Date", "Author", "Description"], rows,
                   widths_cm=[2, 2.8, 3.2, 8.3], caption=None)

    def toc(self, title: str = "Contents", levels: str = "1-3") -> None:
        p = self.d.add_paragraph()
        r = p.add_run(title)
        r.font.size = Pt(17)
        r.font.bold = True
        r.font.color.rgb = ACCENT
        p = self.d.add_paragraph()
        _add_field(p.add_run(), f'TOC \\o "{levels}" \\h \\z \\u')
        self.page_break()

    def list_of(self, kind: str) -> None:
        """List of Tables / Figures (caption-label based TOC field)."""
        p = self.d.add_paragraph()
        r = p.add_run(f"List of {kind}s")
        r.font.size = Pt(17)
        r.font.bold = True
        r.font.color.rgb = ACCENT
        p = self.d.add_paragraph()
        _add_field(p.add_run(), f'TOC \\h \\z \\c "{kind}"')
        self.page_break()

    # ------------------------------------------------------------------ text

    def chapter(self, prefix: str) -> None:
        """Restarts table/figure numbering as <prefix>.n (for dissertations)."""
        self.chapter_prefix = prefix
        self._chapter_tables = 0
        self._chapter_figures = 0

    def h(self, level: int, text: str, numbered: bool = True):
        return self.d.add_heading(text, level=level)

    def p(self, text: str = "", *, align: str | None = None, italic: bool = False,
          size: float | None = None, color: RGBColor | None = None, space_after: float | None = None,
          indent_cm: float | None = None, keep_next: bool = False):
        para = self.d.add_paragraph()
        self._fill(para, text, italic=italic, size=size, color=color)
        if align == "center":
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == "right":
            para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        elif align == "justify":
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if space_after is not None:
            para.paragraph_format.space_after = Pt(space_after)
        if indent_cm is not None:
            para.paragraph_format.left_indent = Cm(indent_cm)
        if keep_next:
            para.paragraph_format.keep_with_next = True
        return para

    def bullets(self, items: Iterable[str], level: int = 1) -> None:
        style = "List Bullet" if level == 1 else "List Bullet 2"
        for item in items:
            para = self.d.add_paragraph(style=style)
            self._fill(para, item)
            para.paragraph_format.space_after = Pt(2)

    def steps(self, items: Iterable[str], start: int = 1) -> None:
        """Numbered list with explicit numbers (restarts per call, unlike List Number)."""
        for i, item in enumerate(items, start):
            para = self.d.add_paragraph()
            para.paragraph_format.left_indent = Cm(0.9)
            para.paragraph_format.first_line_indent = Cm(-0.6)
            para.paragraph_format.space_after = Pt(2)
            self._fill(para, f"{i}.\t{item}")
            para.paragraph_format.tab_stops.add_tab_stop(Cm(0.9))

    def definition(self, term: str, text: str) -> None:
        para = self.d.add_paragraph()
        para.paragraph_format.left_indent = Cm(0.6)
        self._fill(para, f"**{term}** — {text}")

    def note(self, text: str, label: str = "Note") -> None:
        para = self.d.add_paragraph()
        _set_paragraph_fill(para, NOTE_FILL)
        para.paragraph_format.left_indent = Cm(0.3)
        para.paragraph_format.right_indent = Cm(0.3)
        self._fill(para, f"**{label}.** {text}")

    def code(self, text: str, size: float = 8.5) -> None:
        para = self.d.add_paragraph()
        _set_paragraph_fill(para, CODE_FILL)
        para.paragraph_format.space_after = Pt(8)
        para.paragraph_format.line_spacing = 1.0
        lines = text.strip("\n").split("\n")
        for i, line in enumerate(lines):
            run = para.add_run(line)
            run.font.name = "Consolas"
            run._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Consolas")
            run.font.size = Pt(size)
            if i < len(lines) - 1:
                run.add_break()

    def equation(self, text: str, number: str | None = None) -> None:
        """A display line for a formula, set in Cambria Math, numbered on the right."""
        para = self.d.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(text)
        run.font.name = "Cambria Math"
        run.font.size = Pt(11)
        if number:
            run = para.add_run(f"\t({number})")
            para.paragraph_format.tab_stops.add_tab_stop(Cm(16.0))

    def page_break(self) -> None:
        self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def spacer(self, pts: float = 6) -> None:
        para = self.d.add_paragraph()
        para.paragraph_format.space_after = Pt(pts)

    # ----------------------------------------------------------------- tables

    def _next_table_label(self) -> str:
        if self.chapter_prefix:
            self._chapter_tables += 1
            return f"{self.chapter_prefix}{self._chapter_tables}"
        self._table_no += 1
        return str(self._table_no)

    def _next_figure_label(self) -> str:
        if self.chapter_prefix:
            self._chapter_figures += 1
            return f"{self.chapter_prefix}{self._chapter_figures}"
        self._figure_no += 1
        return str(self._figure_no)

    def _caption(self, kind: str, label: str, text: str) -> None:
        para = self.d.add_paragraph(style="Caption")
        para.paragraph_format.keep_with_next = kind == "Table"
        run = para.add_run(f"{kind} ")
        # SEQ field so Word's "List of Tables/Figures" can find the caption;
        # the cached result carries our own numbering.
        seq = para.add_run()
        begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve")
        instr.text = f" SEQ {kind} \\* ARABIC "
        sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
        t = OxmlElement("w:t"); t.text = label
        end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
        for el in (begin, instr, sep, t, end):
            seq._r.append(el)
        para.add_run(f". {text}")

    def table(self, headers: Sequence[str], rows: Sequence[Sequence[str]], *,
              widths_cm: Sequence[float] | None = None, caption: str | None = "",
              font_size: float = 9, header_fill: str = HEADER_FILL,
              first_col_bold: bool = False, zebra: bool = False) -> str | None:
        """Adds a captioned table. Returns the table label ("3", "4.2") or None."""
        label = None
        if caption:
            label = self._next_table_label()
            self._caption("Table", label, caption)
        t = self.d.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        if widths_cm is None:
            usable = 16.3
            widths_cm = [usable / len(headers)] * len(headers)
        hdr = t.rows[0]
        _repeat_header(hdr)
        for i, text in enumerate(headers):
            cell = hdr.cells[i]
            cell.width = Cm(widths_cm[i])
            _set_cell_fill(cell, header_fill)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            self._fill(cell.paragraphs[0], text, bold=True, size=font_size)
        for r_i, row in enumerate(rows):
            cells = t.add_row().cells
            _keep_row_together(t.rows[-1])
            for i, text in enumerate(row):
                cell = cells[i]
                cell.width = Cm(widths_cm[i])
                if zebra and r_i % 2 == 1:
                    _set_cell_fill(cell, "F7F8FA")
                parts = str(text).split("\n")
                self._fill(cell.paragraphs[0], parts[0], size=font_size,
                           bold=first_col_bold and i == 0)
                for extra in parts[1:]:
                    self._fill(cell.add_paragraph(), extra, size=font_size)
                for para in cell.paragraphs:
                    para.paragraph_format.space_after = Pt(1)
                    para.paragraph_format.line_spacing = 1.05
        self.spacer(4)
        return label

    def kv_table(self, rows: Sequence[tuple[str, str]], caption: str | None = None,
                 widths_cm: Sequence[float] = (4.2, 12.1), font_size: float = 9) -> str | None:
        """Two-column field/value table with a shaded first column."""
        label = None
        if caption:
            label = self._next_table_label()
            self._caption("Table", label, caption)
        t = self.d.add_table(rows=0, cols=2)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for k, v in rows:
            cells = t.add_row().cells
            _keep_row_together(t.rows[-1])
            cells[0].width, cells[1].width = Cm(widths_cm[0]), Cm(widths_cm[1])
            _set_cell_fill(cells[0], HEADER_FILL)
            self._fill(cells[0].paragraphs[0], k, bold=True, size=font_size)
            parts = str(v).split("\n")
            self._fill(cells[1].paragraphs[0], parts[0], size=font_size)
            for extra in parts[1:]:
                self._fill(cells[1].add_paragraph(), extra, size=font_size)
            for c in cells:
                for para in c.paragraphs:
                    para.paragraph_format.space_after = Pt(1)
        self.spacer(4)
        return label

    # ---------------------------------------------------------------- figures

    def figure(self, path: str | Path, caption: str, width_cm: float = 14.0) -> str:
        para = self.d.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.keep_with_next = True
        para.add_run().add_picture(str(path), width=Cm(width_cm))
        label = self._next_figure_label()
        self._caption("Figure", label, caption)
        cap = self.d.paragraphs[-1]
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return label

    # --------------------------------------------------------------- plumbing

    @staticmethod
    def _bottom_border(paragraph, color: str, size: int) -> None:
        p_pr = paragraph._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), str(size))
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), color)
        bdr.append(bottom)
        p_pr.append(bdr)

    def _fill(self, paragraph, text: str, *, bold: bool = False, italic: bool = False,
              size: float | None = None, color: RGBColor | None = None) -> None:
        for piece in _INLINE.split(text or ""):
            if not piece:
                continue
            run_bold, run_italic, mono = bold, italic, False
            if piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
                piece, run_bold = piece[2:-2], True
            elif piece.startswith("`") and piece.endswith("`") and len(piece) > 2:
                piece, mono = piece[1:-1], True
            elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
                piece, run_italic = piece[1:-1], True
            run = paragraph.add_run(piece)
            run.bold = run_bold or None
            run.italic = run_italic or None
            if mono:
                run.font.name = "Consolas"
                run._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Consolas")
                run.font.size = Pt((size or 10) - 0.5)
            elif size:
                run.font.size = Pt(size)
            if color is not None:
                run.font.color.rgb = color

    def save(self, path: str | Path) -> Path:
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        self.d.save(str(path))
        return path
